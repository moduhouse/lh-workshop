"""
regression base v9.xlsx -> edat_data.json

원본 엑셀(Sheet1, 9840행 x 45열)을 읽어 대시보드(EDAT_기후적응역량_대시보드.html)가
쓰는 데이터 구조로 변환한다. 컬럼 분리/번들러 없이, 결과 JSON을 HTML 템플릿의
`/*__DATA__*/` 자리에 그대로 붙여넣는 방식으로 쓰인다 (CLAUDE.md "재현 방법" 참고).

실행:
    python scripts/extract_data.py
    (상위 폴더의 "regression base v9.xlsx"를 읽어 같은 폴더에 edat_data.json 생성)
"""
import json
import os
import statistics
from collections import defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_XLSX = os.path.join(ROOT, "regression base v9.xlsx")
OUT_JSON = os.path.join(ROOT, "edat_data.json")

# 국가-연도 집계(cy_rows)에 쓰는 29개 지표.
# slums_pop_Province_* 4개 컬럼(엑셀 12~15열)은 한글 헤더 인코딩이 깨져 있어 제외.
INDICATORS = [
    "water", "elec", "life_both", "life_female", "life_male", "health_exp", "gini",
    "unemploy", "wealth_inequal10", "welath_inequal50", "political", "HDI",
    "Corruption_CPI", "migration", "storm", "flood100yr", "temp", "preci", "urban",
    "urban_pop", "pop_den", "gdppc", "income50", "HappinessScore", "forest", "Green",
    "UHI", "pm2_5", "slums_pop_country",
]

# 지역(province) 상세 테이블(p_rows)에 쓰는 축소 지표 세트 (파일 용량 절감용).
DETAIL_INDICATORS = [
    "temp", "preci", "storm", "flood100yr", "HDI", "gdppc", "urban", "forest",
    "Green", "UHI", "pm2_5", "life_both", "water", "elec", "gini", "unemploy",
    "HappinessScore",
]


def load_rows():
    wb = openpyxl.load_workbook(SRC_XLSX, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=45, values_only=True))
    header, data = rows[0], rows[1:]
    idx = {h: i for i, h in enumerate(header) if h}
    return idx, data


def build_country_year(idx, data):
    agg = defaultdict(lambda: defaultdict(list))
    meta = {}
    for r in data:
        country, year = r[idx["country"]], r[idx["year"]]
        key = (country, year)
        meta[key] = (r[idx["code"]], r[idx["continent"]])
        for ind in INDICATORS:
            v = r[idx[ind]]
            if v is not None:
                agg[key][ind].append(v)

    cols = ["country", "year", "code", "continent"] + INDICATORS
    rows = []
    for (country, year), inds in sorted(agg.items()):
        code, continent = meta[(country, year)]
        row = [country, year, code, continent]
        for ind in INDICATORS:
            vals = inds.get(ind, [])
            row.append(round(statistics.fmean(vals), 2) if vals else None)
        rows.append(row)
    return cols, rows


def build_province_detail(idx, data):
    cols = ["country", "continent", "code", "province", "year"] + DETAIL_INDICATORS
    rows = []
    for r in data:
        row = [r[idx["country"]], r[idx["continent"]], r[idx["code"]], r[idx["province"]], r[idx["year"]]]
        for ind in DETAIL_INDICATORS:
            v = r[idx[ind]]
            row.append(round(v, 2) if isinstance(v, (int, float)) else v)
        rows.append(row)
    return cols, rows


def main():
    idx, data = load_rows()
    cy_cols, cy_rows = build_country_year(idx, data)
    p_cols, p_rows = build_province_detail(idx, data)

    out = {
        "indicators": INDICATORS,
        "cy_cols": cy_cols, "cy_rows": cy_rows,
        "detail_inds": DETAIL_INDICATORS, "p_cols": p_cols, "p_rows": p_rows,
    }
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))

    print(f"국가-연도 행: {len(cy_rows)}개, 지역 상세 행: {len(p_rows)}개")
    print(f"저장 위치: {OUT_JSON} ({os.path.getsize(OUT_JSON) / 1e6:.2f} MB)")
    print("다음: HTML 템플릿의 `const EDAT_DATA = /*__DATA__*/;` 자리에 이 JSON 내용을 그대로 삽입하세요.")


if __name__ == "__main__":
    main()
